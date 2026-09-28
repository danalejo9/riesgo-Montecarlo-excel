#!/usr/bin/env python3
"""
Genera el libro de demostración: el flujo de caja de una empresa ficticia, con
todo preparado para marcar entradas y salidas y simular.

    python3 ejemplos/generar-demo.py

Existe por una razón concreta: probar el complemento sobre un libro real del
usuario mezcla dos incógnitas —¿falla el complemento o falla el modelo?— y
además obliga a tocar un archivo que le importa. Con este libro las dos partes
sabemos exactamente qué hay dentro, qué celdas se marcan y qué resultado
debería salir, así que cualquier diferencia es información sobre el
complemento y no sobre la hoja.

**El modelo es deliberadamente corriente**: una empresa que vende un producto,
con precio, unidades, crecimiento, costes variables y fijos, inversión inicial,
amortización lineal, impuestos y valor residual. Sin una sola función rara. Si
el complemento no puede con esto, no puede con nada.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ANIOS = 5

AZUL = Font(color="1F4E79", bold=True)
TITULO = Font(size=14, bold=True, color="1F4E79")
CABECERA = Font(bold=True, color="FFFFFF")
FONDO_CABECERA = PatternFill("solid", fgColor="1F4E79")
TOTAL = Font(bold=True)
BORDE_ARRIBA = Border(top=Side(style="thin"))
MILES = '#,##0'
EUROS = '#,##0 €'
PORCENTAJE = '0,0%'

wb = Workbook()

# =========================================================================
# Hoja 1: las instrucciones. Va primera para que sea lo que se ve al abrir.
# =========================================================================
ins = wb.active
ins.title = "Instrucciones"

LINEAS = [
    ("t", "Libro de demostración — Riesgo Monte Carlo"),
    ("", ""),
    ("p", "Este libro sirve para probar el complemento sobre un modelo que las dos partes"),
    ("p", "conocemos. Todo está en la hoja «Modelo». No hay nada que rellenar: solo marcar"),
    ("p", "celdas en el panel y pulsar Simular."),
    ("", ""),
    ("h", "1. Marca estas celdas como ENTRADA"),
    ("p", "Selecciona la celda en la hoja «Modelo» y pulsa «Añadir la selección como entrada»."),
    ("p", "El panel propone una triangular alrededor del valor que ya tiene la celda; la"),
    ("p", "columna «distribución sugerida» dice con qué sustituirla."),
    ("", ""),
    ("c", ("Celda", "Qué es", "Valor", "Distribución sugerida")),
    ("e", ("Modelo!C5", "Precio unitario", "120 €", "Triangular  min 105  moda 120  max 140")),
    ("e", ("Modelo!C6", "Unidades vendidas el año 1", "20.000", "Triangular  min 16.000  moda 20.000  max 23.000")),
    ("e", ("Modelo!C7", "Crecimiento anual de unidades", "6 %", "Normal  media 0,06  desviación 0,03")),
    ("e", ("Modelo!C8", "Coste variable unitario", "72 €", "Triangular  min 66  moda 72  max 84")),
    ("e", ("Modelo!C9", "Costes fijos anuales", "480.000 €", "Triangular  min 455.000  moda 480.000  max 530.000")),
    ("", ""),
    ("h", "2. Marca estas celdas como SALIDA"),
    ("p", "Con «Añadir la selección como salida». Son lo que se quiere medir."),
    ("", ""),
    ("c", ("Celda", "Qué es", "Valor con los supuestos base", "")),
    ("e", ("Modelo!C32", "VAN al 11 %", "366.372 €", "")),
    ("e", ("Modelo!C33", "TIR del proyecto", "22,1 %", "")),
    ("e", ("Modelo!C34", "Flujo de caja acumulado a 5 años", "942.185 €", "")),
    ("", ""),
    ("p", "Si tu Excel no da esos tres números al abrir el libro, avísame antes de seguir:"),
    ("p", "significa que el modelo no se está calculando igual, y todo lo demás sobraría."),
    ("", ""),
    ("h", "3. Simula"),
    ("p", "Deja 10.000 iteraciones y el motor en «Automático». Debería tardar menos de un"),
    ("p", "segundo y decir «Rápido (en memoria)»: este libro no usa ninguna función que el"),
    ("p", "motor no conozca."),
    ("", ""),
    ("h", "Qué deberías ver"),
    ("p", "Con las distribuciones de arriba, el proyecto base es bueno pero no seguro:"),
    ("p", "· VAN: mediana en torno a +255.000 €, con un p5 cerca de -660.000 € y un p95"),
    ("p", "  cerca de +1.180.000 €. Aproximadamente un 30 % de probabilidad de VAN negativo."),
    ("p", "· TIR: media en torno al 17 %, frente al 11 % de la tasa de descuento."),
    ("p", "· Un puñado de iteraciones (unas 3 de cada 1.000) saldrán SIN TIR: son escenarios"),
    ("p", "  tan malos que el flujo nunca llega a cambiar de signo, y entonces la TIR no"),
    ("p", "  existe. El panel tiene que decírtelo con un número, no dar «sin datos»."),
    ("", ""),
    ("p", "Con la misma semilla, las mismas celdas y las mismas distribuciones, los números"),
    ("p", "salen idénticos en cualquier equipo y aunque vuelvas a marcar las celdas. Si"),
    ("p", "cambias una distribución o el orden de muestreo, cambian: es otro experimento."),
    ("", ""),
    ("h", "Qué mirar en los resultados"),
    ("p", "· La probabilidad de que el VAN salga negativo. En la curva, la altura donde"),
    ("p", "  cruza el cero: ese porcentaje es el riesgo de que el proyecto destruya valor."),
    ("p", "· El p5 del VAN: el mal escenario razonable, no el peor imaginable."),
    ("p", "· La TIR media frente al 11 % de la tasa de descuento."),
    ("", ""),
    ("h", "Después, para forzar la máquina"),
    ("p", "· Sube a 1.000.000 de iteraciones y compara el tiempo."),
    ("p", "· Cambia el motor a «Compatible (Excel en el bucle)» y baja a 1.000 iteraciones."),
    ("p", "  Ese modo mete a Excel dentro del bucle: es lentísimo a propósito, y lo que"),
    ("p", "  interesa medir es cuántas iteraciones por segundo da tu máquina."),
    ("p", "· Marca también C10 (tasa de descuento) como entrada y mira cuánto se ensancha"),
    ("p", "  la distribución del VAN."),
    ("", ""),
    ("h", "El resaltado"),
    ("p", "Al marcar una celda se pinta en la hoja: las entradas en amarillo y las salidas"),
    ("p", "en verde, el convenio de siempre en modelización financiera. Con «Quitar el"),
    ("p", "resaltado» cada celda vuelve al relleno que tenía antes, incluido «sin relleno»,"),
    ("p", "que no es lo mismo que blanco."),
]

fila = 1
for tipo, contenido in LINEAS:
    if tipo == "t":
        ins.cell(fila, 1, contenido).font = TITULO
    elif tipo == "h":
        ins.cell(fila, 1, contenido).font = AZUL
    elif tipo == "p":
        ins.cell(fila, 1, contenido)
    elif tipo == "c":
        for j, texto in enumerate(contenido):
            celda = ins.cell(fila, 1 + j, texto)
            celda.font = CABECERA
            celda.fill = FONDO_CABECERA
    elif tipo == "e":
        for j, texto in enumerate(contenido):
            ins.cell(fila, 1 + j, texto)
    fila += 1

for col, ancho in [("A", 42), ("B", 34), ("C", 26), ("D", 46)]:
    ins.column_dimensions[col].width = ancho

# =========================================================================
# Hoja 2: el modelo.
# =========================================================================
m = wb.create_sheet("Modelo")

m["B2"] = "Proyecto Aurora — evaluación a cinco años"
m["B2"].font = TITULO
m["B3"] = "Empresa ficticia. Los supuestos en azul son los que conviene tratar como inciertos."
m["B3"].font = Font(italic=True, color="666666")

# --- Supuestos ---
m["B4"] = "SUPUESTOS"
m["B4"].font = AZUL

SUPUESTOS = [
    (5,  "Precio unitario", 120, EUROS),
    (6,  "Unidades vendidas el año 1", 20000, MILES),
    (7,  "Crecimiento anual de unidades", 0.06, PORCENTAJE),
    (8,  "Coste variable unitario", 72, EUROS),
    (9,  "Costes fijos anuales", 480000, EUROS),
    (10, "Tasa de descuento", 0.11, PORCENTAJE),
    (11, "Inflación de costes", 0.03, PORCENTAJE),
    (12, "Inversión inicial", 1200000, EUROS),
    (13, "Tipo impositivo", 0.30, PORCENTAJE),
    (14, "Valor residual en el año 5", 150000, EUROS),
]
for fila, etiqueta, valor, formato in SUPUESTOS:
    m.cell(fila, 2, etiqueta)
    c = m.cell(fila, 3, valor)
    c.number_format = formato
    c.font = Font(color="1F4E79")  # convenio: los datos de entrada, en azul

# --- Flujo de caja ---
m["B17"] = "FLUJO DE CAJA"
m["B17"].font = AZUL

# Cabecera de años: C=año 0, D..H = años 1..5
m.cell(18, 2, "Año").font = TOTAL
for a in range(0, ANIOS + 1):
    c = m.cell(18, 3 + a, a)
    c.font = CABECERA
    c.fill = FONDO_CABECERA
    c.alignment = Alignment(horizontal="center")

def col(a: int) -> str:
    """Letra de la columna del año `a` (el año 0 está en la C)."""
    return get_column_letter(3 + a)

FILAS = {
    "unidades": 19,
    "ingresos": 20,
    "variable": 21,
    "fijos": 22,
    "ebitda": 23,
    "amortizacion": 24,
    "baii": 25,
    "impuestos": 26,
    "neto": 27,
    "flujo": 28,
    "acumulado": 29,
}

ETIQUETAS = [
    (19, "Unidades vendidas", MILES),
    (20, "Ingresos", EUROS),
    (21, "Coste variable", EUROS),
    (22, "Costes fijos", EUROS),
    (23, "EBITDA", EUROS),
    (24, "Amortización", EUROS),
    (25, "Beneficio antes de impuestos", EUROS),
    (26, "Impuestos", EUROS),
    (27, "Beneficio neto", EUROS),
    (28, "Flujo de caja libre", EUROS),
    (29, "Flujo acumulado", EUROS),
]
for fila, etiqueta, _formato in ETIQUETAS:
    m.cell(fila, 2, etiqueta)

for a in range(0, ANIOS + 1):
    C = col(a)
    anterior = col(a - 1) if a > 0 else None

    if a == 0:
        # El año 0 es solo la inversión: no hay operación todavía.
        m[f"{C}28"] = "=-$C$12"
        m[f"{C}29"] = f"={C}28"
    else:
        m[f"{C}19"] = "=$C$6" if a == 1 else f"={anterior}19*(1+$C$7)"
        m[f"{C}20"] = f"={C}19*$C$5"
        # Los costes suben con la inflación desde el año 2.
        m[f"{C}21"] = f"={C}19*$C$8*(1+$C$11)^{a - 1}"
        m[f"{C}22"] = f"=$C$9*(1+$C$11)^{a - 1}"
        m[f"{C}23"] = f"={C}20-{C}21-{C}22"
        m[f"{C}24"] = f"=$C$12/{ANIOS}"
        m[f"{C}25"] = f"={C}23-{C}24"
        # Sin crédito fiscal: un año en pérdidas no paga impuestos, pero
        # tampoco los recupera. Es lo más simple y lo más frecuente.
        m[f"{C}26"] = f"=MAX(0,{C}25*$C$13)"
        m[f"{C}27"] = f"={C}25-{C}26"
        # El flujo libre devuelve la amortización, que no es una salida de caja.
        # El último año se cobra además el valor residual.
        residual = "+$C$14" if a == ANIOS else ""
        m[f"{C}28"] = f"={C}27+{C}24{residual}"
        m[f"{C}29"] = f"={anterior}29+{C}28"

    for fila, _etiqueta, formato in ETIQUETAS:
        celda = m.cell(fila, 3 + a)
        celda.number_format = formato
        if fila in (FILAS["flujo"], FILAS["acumulado"]):
            celda.font = TOTAL
    m.cell(FILAS["flujo"], 3 + a).border = BORDE_ARRIBA

# --- Resultados ---
m["B31"] = "RESULTADOS"
m["B31"].font = AZUL
# Ojo con el convenio de VNA: descuenta desde el primer periodo, así que la
# inversión del año 0 se suma aparte y no entra en el rango.
RESULTADOS = [
    (32, "VAN al 11 %", f"=NPV($C$10,{col(1)}28:{col(ANIOS)}28)+{col(0)}28", EUROS),
    (33, "TIR del proyecto", f"=IRR({col(0)}28:{col(ANIOS)}28)", PORCENTAJE),
    (34, "Flujo de caja acumulado a 5 años", f"={col(ANIOS)}29", EUROS),
    (35, "Margen EBITDA del año 5", f"={col(ANIOS)}23/{col(ANIOS)}20", PORCENTAJE),
]
for fila, etiqueta, formula, formato in RESULTADOS:
    m.cell(fila, 2, etiqueta).font = TOTAL
    c = m.cell(fila, 3, formula)
    c.number_format = formato
    c.font = TOTAL

m["B37"] = "Las celdas a marcar están en la hoja «Instrucciones»."
m["B37"].font = Font(italic=True, color="666666")

m.column_dimensions["A"].width = 3
m.column_dimensions["B"].width = 34
for a in range(0, ANIOS + 1):
    m.column_dimensions[col(a)].width = 14

# =========================================================================
# Hoja 3: pruebas de lo nuevo (funciones RIESGO, LET, matrices, entradas con
# fórmula, ajuste a datos). Hay celdas preparadas y celdas VACÍAS donde el
# usuario escribe la fórmula siguiendo docs/GUIA_PRUEBAS_EXCEL.md: así se
# prueba además que su Excel en español acepta los nombres traducidos.
# =========================================================================
p = wb.create_sheet("Pruebas")
p["B2"] = "Hoja de pruebas — funciones nuevas"
p["B2"].font = TITULO
p["B3"] = "Las celdas en amarillo claro están vacías a propósito: escribe en ellas la fórmula que indica la guía."
p["B3"].font = Font(italic=True, color="666666")
PARA_ESCRIBIR = PatternFill("solid", fgColor="FFF9E6")

def etiqueta(fila, texto, negrita=False, instruccion=None):
    """La etiqueta va sola en B (el panel la usa como nombre de la celda) y la
    instrucción, si hay, en D: antes iban juntas y la entrada acababa llamándose
    «Demanda — escribe =RIESGO.NORMAL(1000;150)»."""
    c = p.cell(fila, 2, texto)
    if negrita:
        c.font = AZUL
    if instruccion:
        d = p.cell(fila, 4, instruccion)
        d.font = Font(italic=True, color="666666")

# 1. Funciones RIESGO
etiqueta(4, "1. FUNCIONES RIESGO (pruebas F1–F6)", True)
etiqueta(5, "Demanda", instruccion="escribe =RIESGO.NORMAL(1000;150)")
etiqueta(6, "Precio", instruccion="escribe =RIESGO.TRIANGULAR(8;10;14)")
etiqueta(7, "Ingresos = Demanda × Precio")
for f in (5, 6):
    p.cell(f, 3).fill = PARA_ESCRIBIR
p["C7"] = "=C5*C6"
p["C7"].number_format = '#,##0.00'

# 2. LET
etiqueta(9, "2. LET (prueba G1)", True)
etiqueta(10, "Margen = ingresos − coste (60 %), con LET")
p["C10"] = "=_xlfn.LET(_xlpm.coste,C7*0.6,C7-_xlpm.coste)"
p["C10"].number_format = '#,##0.00'

# 3. Matrices: ventas por zona, proporcionales a la demanda de C5
etiqueta(12, "3. MATRICES (pruebas G2–G4) — umbral de ventas:", True)
p["C12"] = 100
p["E12"] = "Zona"
p["F12"] = "Ventas"
p["E12"].font = TOTAL
p["F12"].font = TOTAL
for i, (zona, base) in enumerate([("Norte", 120), ("Sur", 80), ("Este", 150), ("Oeste", 60), ("Centro", 200)]):
    p.cell(13 + i, 5, zona)
    c = p.cell(13 + i, 6, f"={base}*$C$5/1000")
    c.number_format = '#,##0.0'
# Las instrucciones de las matrices van en la columna I: en D-F está la tabla.
etiqueta(13, "Ventas por encima del umbral")
etiqueta(14, "Ventas por encima del umbral (SUMAPRODUCTO)")
etiqueta(15, "La mayor venta")
for fila, texto in ((13, "escribe =SUMA(FILTRAR(F13:F17;F13:F17>C12;0))"),
                    (14, "escribe =SUMAPRODUCTO((F13:F17>C12)*F13:F17)"),
                    (15, "escribe =INDICE(ORDENAR(F13:F17;1;-1);1)")):
    c = p.cell(fila, 9, texto)
    c.font = Font(italic=True, color="666666")
for f in (13, 14, 15):
    p.cell(f, 3).fill = PARA_ESCRIBIR
p["H11"] = "Derrame — escribe en H13: =FILTRAR(E13:E17;F13:F17>C12)"
p["H11"].font = TOTAL
p["H13"].fill = PARA_ESCRIBIR

# 4. Entrada con fórmula (modo compatible)
etiqueta(20, "4. ENTRADA CON FÓRMULA (prueba H)", True)
etiqueta(21, "Coste unitario (es una fórmula: =7+1)")
p["C21"] = "=7+1"
etiqueta(22, "Coste total = coste unitario × demanda")
p["C22"] = "=C21*C5"
p["C22"].number_format = '#,##0.00'

# 6. Funciones RIESGO nuevas (B4, I1): discreta, acotada, lognormal por el
#    logaritmo y empírica. Valores y pesos de la discreta en E25:F27.
etiqueta(24, "6. FUNCIONES NUEVAS (pruebas F7–F10)", True)
etiqueta(25, "Plazo de entrega (días)")
etiqueta(26, "Demanda acotada")
etiqueta(27, "Precio lognormal")
etiqueta(28, "Precio según el histórico")
etiqueta(29, "Suma de control")
p["E24"] = "Plazo"
p["F24"] = "Peso"
p["E24"].font = TOTAL
p["F24"].font = TOTAL
for i, (plazo, peso) in enumerate([(3, 0.5), (5, 0.3), (10, 0.2)]):
    p.cell(25 + i, 5, plazo)
    p.cell(25 + i, 6, peso).number_format = '0%'
for fila, texto in ((25, "escribe =RIESGO.DISCRETA(E25:E27;F25:F27)"),
                    (26, "escribe =RIESGO.TRUNCAR(RIESGO.NORMAL(100;30);50;150)"),
                    (27, "escribe =RIESGO.LOGNORMAL2(4,6;0,2)"),
                    (28, "escribe =RIESGO.EMPIRICA(K5:K44)")):
    p.cell(fila, 3).fill = PARA_ESCRIBIR
    c = p.cell(fila, 9, texto)
    c.font = Font(italic=True, color="666666")
p["C29"] = "=C25+C26+C27+C28"
p["C29"].number_format = '#,##0.00'

# 7. RIESGO con celdas como argumentos (B7): fijos, dentro de una operación
#    y con parámetros que dependen de otra celda aleatoria.
etiqueta(31, "7. RIESGO CON CELDAS (pruebas F12–F15)", True)
etiqueta(32, "Media de la demanda")
p["C32"] = 1000
etiqueta(33, "Desviación de la demanda")
p["C33"] = 150
etiqueta(34, "Demanda con celdas")
etiqueta(35, "Precio con imprevisto")
etiqueta(36, "Demanda del mes siguiente")
etiqueta(37, "Ingresos con celdas")
p["C37"] = "=C36*C35"
p["C37"].number_format = '#,##0.00'
for fila, texto in ((34, "escribe =RIESGO.NORMAL(C32;C33)"),
                    (35, "escribe =10+RIESGO.NORMAL(0;1)"),
                    (36, "escribe =RIESGO.NORMAL(C34;50)")):
    p.cell(fila, 3).fill = PARA_ESCRIBIR
    c = p.cell(fila, 9, texto)
    c.font = Font(italic=True, color="666666")

# 8. Resultados de la simulación en la hoja (B6).
etiqueta(39, "8. RESULTADOS EN LA HOJA (pruebas F16–F19)", True)
etiqueta(40, "Media de los ingresos")
etiqueta(41, "p5 de los ingresos")
etiqueta(42, "Probabilidad de ingresos < 8.000")
etiqueta(43, "Desviación de los ingresos")
for fila, texto, formato in ((40, "escribe =RIESGO.MEDIA(C37)", '#,##0.00'),
                             (41, "escribe =RIESGO.PERCENTIL(C37;0,05)", '#,##0.00'),
                             (42, "escribe =RIESGO.PROB(C37;8000)", '0.0%'),
                             (43, "escribe =RIESGO.DESV(C37)", '#,##0.00')):
    c3 = p.cell(fila, 3)
    c3.fill = PARA_ESCRIBIR
    c3.number_format = formato
    c = p.cell(fila, 9, texto)
    c.font = Font(italic=True, color="666666")

# 5. Datos históricos para "Ajustar a datos" (prueba E): 40 precios con cola
#    a la derecha (lognormal de media ~120), deterministas.
p["K4"] = "5. PRECIOS HISTÓRICOS (prueba E)"
p["K4"].font = AZUL
import math
estado = 12345
for i in range(40):
    estado = (estado * 16807) % 2147483647
    u1 = estado / 2147483647
    estado = (estado * 16807) % 2147483647
    u2 = estado / 2147483647
    z = math.sqrt(-2 * math.log(u1)) * math.cos(2 * math.pi * u2)
    p.cell(5 + i, 11, round(118 * math.exp(0.12 * z), 2)).number_format = '#,##0.00'

p.column_dimensions["A"].width = 3
p.column_dimensions["B"].width = 44
p.column_dimensions["D"].width = 12
p.column_dimensions["C"].width = 14
p.column_dimensions["E"].width = 10
p.column_dimensions["F"].width = 10
p.column_dimensions["H"].width = 14
p.column_dimensions["K"].width = 12
p.column_dimensions["I"].width = 50

# =========================================================================
# Hoja 4: variables de texto (bloque T de la guía). Una zona de venta que es
# un texto («Norte», «Sur», «Este») con su probabilidad, fórmulas que la usan
# (BUSCARV), y salidas de texto y de sí/no.
# =========================================================================
z = wb.create_sheet("Zonas")
z["B2"] = "Variables de texto — zona de venta"
z["B2"].font = TITULO
z["B3"] = "La zona es un texto con su probabilidad. El recargo depende de la zona; la gama y el «¿supera?» son salidas de texto y de sí/no."
z["B3"].font = Font(italic=True, color="666666")

z["E4"] = "Zona"
z["F4"] = "Probabilidad"
z["G4"] = "Recargo"
for c in ("E4", "F4", "G4"):
    z[c].font = TOTAL
for i, (zona, prob, recargo) in enumerate([("Norte", 0.2, 1.10), ("Sur", 0.5, 1.00), ("Este", 0.3, 0.90)]):
    z.cell(5 + i, 5, zona)
    z.cell(5 + i, 6, prob).number_format = '0%'
    z.cell(5 + i, 7, recargo).number_format = '0.00'

filas_z = [
    (4, "Zona de venta", "Norte", None, "marca como entrada: el panel la propone como texto"),
    (5, "Recargo por zona", "=VLOOKUP(C4,$E$5:$G$7,3,FALSE)", '0.00', None),
    (6, "Precio base", 100, '#,##0 €', "marca como entrada"),
    (7, "Unidades", 1000, '#,##0', "marca como entrada"),
    (8, "Ingresos", "=C6*C5*C7", '#,##0 €', "salida numérica"),
    (9, "Gama", '=IF(C8>110000,"Alta",IF(C8>90000,"Media","Baja"))', None, "salida de texto"),
    (10, "¿Supera 100.000 €?", "=C8>100000", None, "salida de sí/no"),
    (12, "Zona con función", None, None, "escribe =RIESGO.CATEGORIA(E5:E7;F5:F7)"),
    (13, "Ingresos con la zona de la función", "=C6*VLOOKUP(C12,$E$5:$G$7,3,FALSE)*C7", '#,##0 €', "salida"),
]
for fila, texto, valor, formato, nota in filas_z:
    z.cell(fila, 2, texto)
    if valor is not None:
        c = z.cell(fila, 3, valor)
        if formato:
            c.number_format = formato
    if nota:
        n = z.cell(fila, 4, nota) if fila not in (4, 6, 7, 8, 9, 10, 12, 13) else z.cell(fila, 9, nota)
        n.font = Font(italic=True, color="666666")
z["C12"].fill = PARA_ESCRIBIR
z.column_dimensions["A"].width = 3
z.column_dimensions["B"].width = 36
z.column_dimensions["C"].width = 14
z.column_dimensions["E"].width = 10
z.column_dimensions["F"].width = 12
z.column_dimensions["G"].width = 10
z.column_dimensions["I"].width = 44

wb.calculation.fullCalcOnLoad = True

destino = "ejemplos/demo-proyecto-aurora.xlsx"
wb.save(destino)
print(f"Escrito {destino}")
