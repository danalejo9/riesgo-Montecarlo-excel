#!/usr/bin/env python3
"""
Genera el libro de pruebas del cronograma (frente C, bloque Q de la guía):

    python3 ejemplos/generar-cronograma.py

Es al cronograma lo que `demo-proyecto-aurora.xlsx` es al modelo de costos:
un proyecto que las dos partes conocemos de memoria, con **el resultado
esperado ya escrito** en la hoja «Esperado». Así, cualquier diferencia entre
lo que calcula el complemento y lo que dice esa hoja es información sobre el
complemento, no sobre el cronograma.

Los valores esperados **no** salen del motor del complemento: los calcula
este script con su propio CPM, escrito aparte y en otro lenguaje. Si los dos
coinciden, no es porque compartan el mismo error. `tests/unit/cronogramaDemo.test.ts`
compara el motor con esta hoja en cada `npm test`.

El proyecto (un centro comunitario de dos plantas) es corriente a propósito,
pero toca todo lo que el CPM tiene que saber hacer: las cuatro relaciones
(FC, CC, FF, CF), retardos y un adelanto, hitos, una convergencia de varias
rutas, un «inicio no antes de», festivos y un suministro de plazo largo que
compite con la obra.
"""

import datetime as dt
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INICIO = dt.date(2027, 2, 1)  # lunes
LABORABLES = {0, 1, 2, 3, 4}  # lunes a viernes (weekday() de Python)
FESTIVOS = [
    (dt.date(2027, 3, 25), "Jueves Santo"),
    (dt.date(2027, 3, 26), "Viernes Santo"),
    (dt.date(2027, 5, 1), "Día del Trabajo (cae en sábado)"),
    (dt.date(2027, 7, 20), "Festivo nacional"),
    (dt.date(2027, 8, 16), "Festivo nacional"),
]

# (ID, actividad, duración, predecesoras, mín, más probable, máx, categoría, costo por día, inicio no antes de)
ACTIVIDADES = [
    ("1", "Inicio del proyecto", 0, "", None, None, None, "Hitos", 0, None),
    ("2", "Permisos y licencias", 15, "1", 10, 15, 35, "Trámites", 200, None),
    ("3", "Diseño detallado", 20, "1", 18, 20, 28, "Diseño", 900, None),
    ("4", "Movilización de obra", 5, "2", 4, 5, 8, "Obra gris", 1500, None),
    ("5", "Excavación", 10, "3;4", 8, 10, 16, "Obra gris", 2500, None),
    ("6", "Cimentación", 15, "5", 13, 15, 22, "Obra gris", 3000, None),
    ("7", "Estructura", 30, "6", 27, 30, 42, "Obra gris", 3500, None),
    ("8", "Fabricación del ascensor", 45, "3FC+5", 40, 45, 70, "Suministros", 100, None),
    ("9", "Cubierta", 12, "7", 10, 12, 18, "Obra gris", 2200, None),
    ("10", "Instalaciones eléctricas", 20, "7CC+10", 18, 20, 26, "Instalaciones", 1800, None),
    ("11", "Instalaciones hidráulicas", 18, "7CC+10", 16, 18, 25, "Instalaciones", 1600, None),
    ("12", "Mampostería", 20, "7FC-5", 18, 20, 28, "Obra gris", 2000, None),
    ("13", "Montaje del ascensor", 8, "8;9", 7, 8, 12, "Suministros", 1200, None),
    ("14", "Acabados interiores", 25, "10;11;12", 22, 25, 35, "Acabados", 2400, None),
    ("15", "Fachada", 15, "9", 13, 15, 20, "Acabados", 1900, None),
    ("16", "Pruebas de instalaciones", 6, "10FF+3;11FF+3", 5, 6, 10, "Instalaciones", 800, None),
    ("17", "Paisajismo", 10, "15", 8, 10, 14, "Exteriores", 700, dt.date(2027, 6, 28)),
    ("18", "Limpieza final", 3, "13;14;16;17", 3, 3, 5, "Acabados", 500, None),
    ("19", "Entrega al cliente", 0, "18", None, None, None, "Hitos", 0, None),
    ("20", "Vigilancia hasta la entrega", 10, "19CF", None, None, None, "Obra gris", 150, None),
]

TIPOS = {"FC": "FS", "FS": "FS", "CC": "SS", "SS": "SS", "FF": "FF", "CF": "SF", "SF": "SF"}


# ---------------------------------------------------------------------------
# El CPM de referencia, independiente del motor en TypeScript.
# ---------------------------------------------------------------------------
def dias_laborables(desde, cuantos):
    festivos = {f for f, _ in FESTIVOS}
    dias, d = [], desde
    while len(dias) < cuantos:
        if d.weekday() in LABORABLES and d not in festivos:
            dias.append(d)
        d += dt.timedelta(days=1)
    return dias


DIAS = dias_laborables(INICIO, 2000)


def indice_desde(fecha):
    return next(i for i, d in enumerate(DIAS) if d >= fecha)


def leer_preds(texto):
    rel = []
    for parte in [p.strip() for p in texto.replace(",", ";").split(";") if p.strip()]:
        num = ""
        i = 0
        while i < len(parte) and parte[i].isdigit():
            num += parte[i]
            i += 1
        resto = parte[i:]
        tipo = "FS"
        if resto[:2].upper() in TIPOS:
            tipo = TIPOS[resto[:2].upper()]
            resto = resto[2:]
        retardo = float(resto) if resto else 0.0
        rel.append((num, tipo, retardo))
    return rel


def indice_hasta(fecha):
    return indice_desde(fecha + dt.timedelta(days=1)) - 1


# Escenario con avance (prueba del resto de P1): fecha de corte, fechas reales,
# % completado, duración restante y restricciones de fin. Por actividad:
# (inicio real, fin real, % completado, duración restante,
#  terminar no antes de, terminar no después de)
CORTE = dt.date(2027, 3, 15)
AVANCE = {
    "1": (dt.date(2027, 2, 1), dt.date(2027, 2, 1), None, None, None, None),
    "2": (dt.date(2027, 2, 1), dt.date(2027, 2, 26), None, None, None, None),
    "3": (dt.date(2027, 2, 1), dt.date(2027, 2, 26), None, None, None, None),
    "4": (dt.date(2027, 3, 1), dt.date(2027, 3, 5), None, None, None, None),
    "5": (dt.date(2027, 3, 8), None, 0.4, None, None, None),
    "8": (dt.date(2027, 3, 8), None, None, 42, None, None),
    "15": (None, None, None, None, None, dt.date(2027, 6, 30)),
    "17": (None, None, None, None, dt.date(2027, 7, 21), None),
}


def cpm(avance=None, corte=None):
    """CPM de referencia. Con `avance`, las reglas de engine/cronograma/red.ts."""
    avance = avance or {}
    acts = {a[0]: a for a in ACTIVIDADES}
    preds = {a[0]: leer_preds(a[3]) for a in ACTIVIDADES}
    C = indice_desde(corte) if corte else 0
    dur, min_ini, min_fin, max_fin, sin_logica, estado = {}, {}, {}, {}, {}, {}
    for a in ACTIVIDADES:
        k, d = a[0], a[2]
        ir, fr, comp, rest, fnet, fnlt = avance.get(k, (None,) * 6)
        max_fin[k] = indice_hasta(fnlt) + 1 if fnlt else math.inf
        min_fin[k] = -math.inf
        if fr:
            fin = indice_hasta(fr) + 1
            # Un hito terminado ocupa cero días: empieza donde termina.
            ini = fin if d == 0 else max(0, min(indice_desde(ir) if ir else fin - d, fin))
            estado[k], min_ini[k], sin_logica[k], dur[k] = "Terminada", ini, True, fin - ini
        elif ir:
            ini = indice_desde(ir)
            r = rest if rest is not None else d * (1 - (comp or 0))
            estado[k], min_ini[k], sin_logica[k], dur[k] = "En curso", ini, True, max(C, ini) - ini + r
        else:
            v = C
            if acts[k][9] is not None:
                v = max(v, indice_desde(acts[k][9]))
            estado[k], min_ini[k], sin_logica[k], dur[k] = "Pendiente", v, False, d
            if fnet:
                min_fin[k] = indice_desde(fnet) + 1
    orden, hecho = [], set()
    while len(orden) < len(ACTIVIDADES):  # los IDs no están necesariamente en orden
        for a in ACTIVIDADES:
            if a[0] not in hecho and all(p in hecho for p, _, _ in preds[a[0]]):
                orden.append(a[0])
                hecho.add(a[0])
    es, ef = {}, {}
    for s in orden:
        v = min_ini[s]
        if not sin_logica[s]:
            for p, t, r in preds[s]:
                v = max(v, {"FS": ef[p] + r, "SS": es[p] + r,
                            "FF": ef[p] + r - dur[s], "SF": es[p] + r - dur[s]}[t])
            v = max(v, min_fin[s] - dur[s])
        es[s], ef[s] = v, v + dur[s]
    fin = max(ef.values())
    sucs = {a[0]: [] for a in ACTIVIDADES}
    for s in orden:
        for p, t, r in preds[s]:
            sucs[p].append((s, t, r))
    ls, lf = {}, {}
    for p in reversed(orden):
        v = fin
        for s, t, r in sucs[p]:
            v = min(v, {"FS": ls[s] - r, "SS": ls[s] - r + dur[p],
                        "FF": lf[s] - r, "SF": lf[s] - r + dur[p]}[t])
        v = min(v, max_fin[p])
        lf[p], ls[p] = v, v - dur[p]
    for k in estado:
        if estado[k] == "Terminada":
            ls[k], lf[k] = es[k], ef[k]
    return es, ef, ls, lf, fin, dur, estado


def fecha_fin(f, d):
    return DIAS[max(0, math.ceil(f) - 1)] if (d > 0 or f > 0) else DIAS[0]


def fecha_inicio(i, d):
    return fecha_fin(i, 0) if d == 0 else DIAS[int(math.floor(i))]


ES, EF, LS, LF, FIN, _, ESTADO = cpm()
ES2, EF2, LS2, LF2, FIN2, DUR2, ESTADO2 = cpm(AVANCE, CORTE)

# ---------------------------------------------------------------------------
# Riesgos de plazo (risk drivers, AACE 57R-09): probabilidad, impacto en % de
# la duración y actividades afectadas (por ID o por categoría).
# ---------------------------------------------------------------------------
RIESGOS = [
    ("R1", "Lluvias intensas durante la obra gris", 0.50, 0.05, 0.10, 0.25, "Obra gris"),
    ("R2", "Demora en la licencia de construcción", 0.30, 0.20, 0.40, 0.80, "2"),
    ("R3", "Retraso del proveedor del ascensor", 0.40, 0.10, 0.20, 0.50, "8"),
    ("R4", "Bajo rendimiento de la cuadrilla de acabados", 0.25, 0.05, 0.15, 0.30, "Acabados"),
    ("R5", "Oportunidad: estructura prefabricada", 0.20, -0.15, -0.10, -0.05, "7"),
]
ITERACIONES_MC = 20000

# Costos (P5): costo fijo por actividad (materiales, equipos comprados), costo
# indirecto del proyecto por día laborable y costo directo de algunos riesgos.
COSTO_FIJO = {"5": 18000, "6": 65000, "7": 140000, "8": 95000, "9": 30000, "10": 42000, "11": 36000,
              "12": 25000, "14": 60000, "15": 38000, "17": 12000}
COSTO_INDIRECTO_DIA = 1800
COSTO_RIESGO = {"R2": 12000, "R3": 20000}


def afectadas(texto):
    ids = {a[0] for a in ACTIVIDADES}
    out = []
    for t in [x.strip() for x in texto.split(";") if x.strip()]:
        out += [t] if t in ids else [a[0] for a in ACTIVIDADES if a[7] == t]
    return out


def pert_muestra(rng, a, m, b):
    if b <= a:
        return m
    al = 1 + 4 * (m - a) / (b - a)
    be = 1 + 4 * (b - m) / (b - a)
    return a + (b - a) * rng.betavariate(al, be)


def monte_carlo():
    import random
    rng = random.Random(20270201)
    acts = {a[0]: a for a in ACTIVIDADES}
    preds = {a[0]: leer_preds(a[3]) for a in ACTIVIDADES}
    orden, hecho = [], set()
    while len(orden) < len(ACTIVIDADES):
        for a in ACTIVIDADES:
            if a[0] not in hecho and all(p in hecho for p, _, _ in preds[a[0]]):
                orden.append(a[0]); hecho.add(a[0])
    sucs = {a[0]: [] for a in ACTIVIDADES}
    for s_ in orden:
        for p, t, r in preds[s_]:
            sucs[p].append((s_, t, r))
    ina = {a[0]: (indice_desde(a[9]) if a[9] else 0) for a in ACTIVIDADES}
    riesgos = [(r, afectadas(r[6])) for r in RIESGOS]
    fines, crit, costos, fines_dia = [], {a[0]: 0 for a in ACTIVIDADES}, [], []
    fijo = sum(COSTO_FIJO.values())
    for _ in range(ITERACIONES_MC):
        d = {}
        for a in ACTIVIDADES:
            d[a[0]] = pert_muestra(rng, a[4], a[5], a[6]) if a[4] is not None else a[2]
        costo = fijo
        for (rid, _, p, lo, mp, hi, _), ids in riesgos:
            if rng.random() < p:
                f = 1 + pert_muestra(rng, lo, mp, hi)
                for i in ids:
                    d[i] *= f
                costo += COSTO_RIESGO.get(rid, 0)
        es, ef = {}, {}
        for s_ in orden:
            v = ina[s_]
            for p, t, r in preds[s_]:
                v = max(v, {"FS": ef[p] + r, "SS": es[p] + r, "FF": ef[p] + r - d[s_], "SF": es[p] + r - d[s_]}[t])
            es[s_], ef[s_] = v, v + d[s_]
        fin = max(ef.values())
        lf = {}
        for p in reversed(orden):
            v = fin
            for s_, t, r in sucs[p]:
                v = min(v, {"FS": lf[s_] - d[s_] - r, "SS": lf[s_] - d[s_] - r + d[p],
                            "FF": lf[s_] - r, "SF": lf[s_] - r + d[p]}[t])
            lf[p] = v
            if v - d[p] - es[p] < 1e-9:
                crit[p] += 1
        fines.append(max(0, math.ceil(fin - 1e-9) - 1))  # día de fin (índice laborable)
        costo += sum(acts[k][8] * d[k] for k in d) + COSTO_INDIRECTO_DIA * fin
        costos.append(costo)
    pares = list(zip(fines, costos))
    fines_ord, costos_ord = sorted(fines), sorted(costos)
    def cuantil(lista, p):
        return lista[max(1, math.ceil(p * len(lista) - 1e-9)) - 1]
    d80, c80 = cuantil(fines_ord, 0.8), cuantil(costos_ord, 0.8)
    jcl = sum(1 for f, c in pares if f <= d80 and c <= c80) / len(pares)
    costo_pct = {p: cuantil(costos_ord, p) for p in (0.1, 0.5, 0.8, 0.9)}
    return ({p: cuantil(fines_ord, p) for p in (0.1, 0.5, 0.8, 0.9)}, {k: v / ITERACIONES_MC for k, v in crit.items()},
            costo_pct, jcl)


def salud():
    """El chequeo de salud, calculado aparte (umbrales DCMA)."""
    tareas = [a for a in ACTIVIDADES if a[2] > 0]
    rel = [(a[0], p, t, r) for a in ACTIVIDADES for p, t, r in leer_preds(a[3])]
    con_suc = {p for _, p, _, _ in rel}
    sin_pred = [a[0] for a in ACTIVIDADES if not a[3]]
    sin_suc = [a[0] for a in ACTIVIDADES if a[0] not in con_suc]
    abiertos = len(sin_pred) - 1 + len(sin_suc) - 1
    nT, nR = len(tareas), len(rel)
    no_fc = [x for x in rel if x[2] != "FS"]
    cf = [x for x in rel if x[2] == "SF"]
    hol = {a[0]: LS[a[0]] - ES[a[0]] for a in ACTIVIDADES}
    return [
        ("logica", "verde" if abiertos == 0 else ("rojo" if abiertos / len(ACTIVIDADES) > 0.05 else "amarillo")),
        ("adelantos", "amarillo" if any(r < 0 for *_, r in rel) else "verde"),
        ("retardos", "amarillo" if sum(1 for *_, r in rel if r > 0) / nR > 0.05 else "verde"),
        ("relaciones", "amarillo" if cf or len(no_fc) / nR > 0.1 else "verde"),
        ("restricciones", (lambda k: "amarillo" if k / len(ACTIVIDADES) > 0.05 else ("info" if k else "verde"))(sum(1 for a in ACTIVIDADES if a[9]))),
        ("holgura-negativa", "rojo" if any(h < -1e-9 for h in hol.values()) else "verde"),
        ("holgura-alta", "amarillo" if sum(1 for a in tareas if hol[a[0]] > 44) / nT > 0.05 else "verde"),
        ("duracion-larga", "amarillo" if sum(1 for a in tareas if a[2] > 44) / nT > 0.05 else "verde"),
        ("incertidumbre", (lambda k: "verde" if k == 0 else ("amarillo" if k / nT > 0.5 else "info"))(sum(1 for a in tareas if a[4] is None or a[4] == a[6]))),
        ("convergencias", "info" if any(len(leer_preds(a[3])) >= 3 for a in ACTIVIDADES) else "verde"),
    ]


PCT, CRIT, COSTO_PCT, JCL = monte_carlo()
COSTO_DET = (sum(COSTO_FIJO.values()) + sum(a[8] * a[2] for a in ACTIVIDADES) + COSTO_INDIRECTO_DIA * FIN)
SALUD = salud()

# ---------------------------------------------------------------------------
# El libro.
# ---------------------------------------------------------------------------
AZUL = "1F4E79"
TITULO = Font(size=14, bold=True, color=AZUL)
H = Font(bold=True, color=AZUL)
CAB = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor=AZUL)
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
FECHA = "dd/mm/yyyy"
DINERO = '"US$" #,##0'

wb = Workbook()
ins = wb.active
ins.title = "Instrucciones"
LINEAS = [
    ("t", "Libro de pruebas del cronograma — Riesgo Monte Carlo"),
    ("", ""),
    ("p", "Es el libro del frente C (módulo Cronograma). Lo generó ejemplos/generar-cronograma.py;"),
    ("p", "no lo edites a mano: si hay que cambiarlo, se cambia el script y se vuelve a generar."),
    ("", ""),
    ("h", "Qué hay en cada hoja"),
    ("p", "«Actividades»: el cronograma de un centro comunitario, 20 actividades, con las cuatro"),
    ("p", "relaciones (FC, CC, FF, CF), retardos, un adelanto, dos hitos, un «inicio no antes de»"),
    ("p", "y la incertidumbre de cada duración (mínimo, más probable, máximo) para el Monte Carlo."),
    ("p", "«Calendario»: inicio del proyecto, días laborables (lunes a viernes) y festivos."),
    ("p", "«Riesgos de plazo»: cinco riesgos (uno es una oportunidad) que alargan a la vez sus actividades."),
    ("p", "«Esperado»: las fechas que debe dar el complemento, calculadas aparte por el script."),
    ("p", "«Esperado (salud)» y «Esperado (simulación)»: el semáforo y el Monte Carlo de referencia."),
    ("p", "«Avance (prueba)» y «Esperado (con avance)»: un escenario con fecha de corte, fechas reales,"),
    ("p", "% completado, duración restante y restricciones de fin, con su resultado esperado."),
    ("", ""),
    ("h", "Cómo se usa"),
    ("p", "Panel → pestaña Cronograma → «Calcular cronograma». Escribe la hoja «Cronograma (resultado)»."),
    ("p", "Compárala con «Esperado»: mismas fechas, misma holgura y mismas actividades críticas."),
    ("p", "Los pasos exactos, y qué tocar para provocar cada error, están en la guía de pruebas,"),
    ("p", "bloque Q (docs/GUIA_PRUEBAS_EXCEL.md)."),
    ("", ""),
    ("h", "Resultado esperado, en una línea"),
    ("p", f"Fin del proyecto: {fecha_fin(FIN, 1).strftime('%d/%m/%Y')} · {FIN:g} días laborables · "
          f"{sum(1 for a in ACTIVIDADES if abs(LS[a[0]] - ES[a[0]]) < 1e-9)} actividades críticas."),
]
for i, (tipo, texto) in enumerate(LINEAS, start=1):
    c = ins.cell(row=i, column=1, value=texto)
    if tipo == "t":
        c.font = TITULO
    elif tipo == "h":
        c.font = H
ins.column_dimensions["A"].width = 100

# Actividades: título en la fila 1, encabezado en la 3 (el lector lo busca).
act = wb.create_sheet("Actividades")
act["A1"] = "Cronograma — Centro comunitario"
act["A1"].font = TITULO
CABECERA = ["ID", "Actividad", "Duración (días)", "Predecesoras", "Mínimo", "Más probable",
            "Máximo", "Categoría", "Costo por día", "Inicio no antes de",
            # Avance y más restricciones (vacías en el escenario base; ver «Avance (prueba)»).
            "Inicio real", "Fin real", "% completado", "Duración restante",
            "Inicio no después de", "Terminar no antes de", "Terminar no después de", "Costo fijo"]
for j, t in enumerate(CABECERA, start=1):
    c = act.cell(row=3, column=j, value=t)
    c.font, c.fill = CAB, FONDO
for i, a in enumerate(ACTIVIDADES, start=4):
    for j, v in enumerate(a, start=1):
        c = act.cell(row=i, column=j, value=v if v != "" else None)
        if j == 9:
            c.number_format = DINERO
        if j == 10 and v is not None:
            c.number_format = FECHA
    act.cell(row=i, column=4).alignment = Alignment(horizontal="left")
for i, a in enumerate(ACTIVIDADES, start=4):
    if a[0] in COSTO_FIJO:
        act.cell(row=i, column=18, value=COSTO_FIJO[a[0]]).number_format = DINERO
for j, w in enumerate([6, 32, 15, 16, 10, 13, 10, 15, 14, 18, 12, 12, 13, 17, 19, 19, 21, 13], start=1):
    act.column_dimensions[get_column_letter(j)].width = w
for i in range(4, 4 + len(ACTIVIDADES)):
    for j in (11, 12, 15, 16, 17):
        act.cell(row=i, column=j).number_format = FECHA
    act.cell(row=i, column=13).number_format = "0%"
act.freeze_panes = "C4"

cal = wb.create_sheet("Calendario")
cal["A1"], cal["B1"] = "Inicio del proyecto", INICIO
cal["B1"].number_format = FECHA
cal["A1"].font = H
cal["A2"] = "Fecha de corte"
cal["B2"].number_format = FECHA
cal["C2"] = "(vacía: sin avance; ver «Avance (prueba)»)"
cal["A3"] = "Días laborables"
cal["A3"].font = H
for k, (dia, lab) in enumerate([("Lunes", 1), ("Martes", 1), ("Miércoles", 1), ("Jueves", 1),
                                  ("Viernes", 1), ("Sábado", 0), ("Domingo", 0)], start=4):
    cal.cell(row=k, column=1, value=dia)
    cal.cell(row=k, column=2, value="Sí" if lab else "No")
cal["A11"], cal["B11"] = "Costo indirecto por día", COSTO_INDIRECTO_DIA
cal["B11"].number_format = DINERO
cal["A11"].font = H
cal["A12"] = "Festivos"
cal["A12"].font = H
for k, (f, nombre) in enumerate(FESTIVOS, start=13):
    cal.cell(row=k, column=1, value=f).number_format = FECHA
    cal.cell(row=k, column=2, value=nombre)
cal.column_dimensions["A"].width = 20
cal.column_dimensions["B"].width = 34

def hoja_esperado(nombre, titulo, es, ef, ls, lf, fin, dur, estado):
    h = wb.create_sheet(nombre)
    h["A1"] = titulo
    h["A1"].font = H
    h["A2"], h["B2"] = "Fin del proyecto", fecha_fin(fin, 1)
    h["B2"].number_format = FECHA
    h["A3"], h["B3"] = "Duración (días laborables)", fin
    cab = ["ID", "Actividad", "Duración", "Inicio", "Fin", "Inicio tardío", "Fin tardío",
           "Holgura total", "Crítica", "Estado"]
    for j, t in enumerate(cab, start=1):
        c = h.cell(row=5, column=j, value=t)
        c.font, c.fill = CAB, FONDO
    for i, a in enumerate(ACTIVIDADES, start=6):
        k, d = a[0], dur[a[0]]
        holgura = 0 if estado[k] == "Terminada" else ls[k] - es[k]
        critica = estado[k] != "Terminada" and holgura < 1e-9
        fila = [k, a[1], a[2], fecha_inicio(es[k], d), fecha_fin(ef[k], d),
                fecha_inicio(max(0, ls[k]), d), fecha_fin(lf[k], d),
                0 if abs(holgura) < 1e-9 else holgura, "Sí" if critica else "No", estado[k]]
        for j, v in enumerate(fila, start=1):
            c = h.cell(row=i, column=j, value=v)
            if 4 <= j <= 7:
                c.number_format = FECHA
            if j == 9 and v == "Sí":
                c.fill = AMARILLO
    for j, w in enumerate([6, 32, 10, 12, 12, 13, 12, 13, 8, 11], start=1):
        h.column_dimensions[get_column_letter(j)].width = w


hoja_esperado("Esperado", "Resultado esperado del CPM determinista (calculado por el script, no por el complemento)",
              ES, EF, LS, LF, FIN, {a[0]: a[2] for a in ACTIVIDADES}, ESTADO)

rie = wb.create_sheet("Riesgos de plazo")
rie["A1"] = "Riesgos de plazo — cada uno alarga a la vez todas sus actividades (AACE 57R-09)"
rie["A1"].font = TITULO
CAB_R = ["ID", "Riesgo", "Probabilidad", "Impacto mínimo", "Impacto más probable", "Impacto máximo",
         "Actividades afectadas", "Costo si ocurre"]
for j, t in enumerate(CAB_R, start=1):
    c = rie.cell(row=3, column=j, value=t)
    c.font, c.fill = CAB, FONDO
for i, r in enumerate(RIESGOS, start=4):
    for j, v in enumerate(r, start=1):
        c = rie.cell(row=i, column=j, value=v)
        if 3 <= j <= 6:
            c.number_format = "0%"
for i, r in enumerate(RIESGOS, start=4):
    if r[0] in COSTO_RIESGO:
        rie.cell(row=i, column=8, value=COSTO_RIESGO[r[0]]).number_format = DINERO
for j, w in enumerate([6, 44, 12, 15, 20, 15, 22, 15], start=1):
    rie.column_dimensions[get_column_letter(j)].width = w

sal = wb.create_sheet("Esperado (salud)")
sal["A1"] = "Chequeo de salud esperado (calculado por el script)"
sal["A1"].font = H
for j, t in enumerate(["Comprobación", "Semáforo"], start=1):
    c = sal.cell(row=3, column=j, value=t)
    c.font, c.fill = CAB, FONDO
for i, (cod, est) in enumerate(SALUD, start=4):
    sal.cell(row=i, column=1, value=cod)
    sal.cell(row=i, column=2, value=est)
sal.column_dimensions["A"].width = 22

sim = wb.create_sheet("Esperado (simulación)")
sim["A1"] = (f"Monte Carlo de referencia: {ITERACIONES_MC} iteraciones en Python (PERT + riesgos de plazo). "
             "El complemento usa otros números aleatorios: debe coincidir en ±1 día, ±3 puntos de criticidad, "
             "±1 % de costo y ±3 puntos de JCL. El costo determinista es exacto.")
sim["A1"].font = H
sim["A3"], sim["B3"], sim["C3"] = "Percentil", "Días laborables", "Fecha de fin"
for c in (sim["A3"], sim["B3"], sim["C3"]):
    c.font, c.fill = CAB, FONDO
for i, p in enumerate((0.1, 0.5, 0.8, 0.9), start=4):
    sim.cell(row=i, column=1, value=f"P{int(p * 100)}")
    sim.cell(row=i, column=2, value=PCT[p] + 1)
    sim.cell(row=i, column=3, value=DIAS[PCT[p]]).number_format = FECHA
sim["E3"], sim["F3"] = "Costo", "Valor"
for c in (sim["E3"], sim["F3"]):
    c.font, c.fill = CAB, FONDO
for i, (nombre, v) in enumerate([("Determinista (exacto)", COSTO_DET), ("P10", COSTO_PCT[0.1]), ("P50", COSTO_PCT[0.5]),
                                  ("P80", COSTO_PCT[0.8]), ("P90", COSTO_PCT[0.9]), ("JCL: P(plazo ≤ P80 y costo ≤ P80)", JCL)],
                                 start=4):
    sim.cell(row=i, column=5, value=nombre)
    sim.cell(row=i, column=6, value=round(v, 4) if nombre.startswith("JCL") else round(v, 2)).number_format = "0%" if nombre.startswith("JCL") else DINERO
sim.column_dimensions["E"].width = 34
sim.column_dimensions["F"].width = 16
sim["A10"], sim["B10"], sim["C10"] = "ID", "Actividad", "Índice de criticidad"
for c in (sim["A10"], sim["B10"], sim["C10"]):
    c.font, c.fill = CAB, FONDO
for i, a in enumerate(ACTIVIDADES, start=11):
    sim.cell(row=i, column=1, value=a[0])
    sim.cell(row=i, column=2, value=a[1])
    sim.cell(row=i, column=3, value=round(CRIT[a[0]], 3)).number_format = "0%"
sim.column_dimensions["A"].width = 12
sim.column_dimensions["B"].width = 32
sim.column_dimensions["C"].width = 18

av = wb.create_sheet("Avance (prueba)")
av["A1"] = "Escenario con avance: copia estos valores en «Actividades» (mismas columnas) y la fecha de corte en «Calendario»!B2"
av["A1"].font = H
av["A2"], av["B2"] = "Fecha de corte", CORTE
av["B2"].number_format = FECHA
CAB_AV = ["ID", "Inicio real", "Fin real", "% completado", "Duración restante", "Terminar no antes de",
          "Terminar no después de"]
for j, t in enumerate(CAB_AV, start=1):
    c = av.cell(row=4, column=j, value=t)
    c.font, c.fill = CAB, FONDO
for i, (k, (ir, fr, comp, rest, fnet, fnlt)) in enumerate(AVANCE.items(), start=5):
    for j, v in enumerate([k, ir, fr, comp, rest, fnet, fnlt], start=1):
        c = av.cell(row=i, column=j, value=v)
        if j in (2, 3, 6, 7):
            c.number_format = FECHA
        if j == 4:
            c.number_format = "0%"
for j, w in enumerate([6, 12, 12, 13, 17, 20, 22], start=1):
    av.column_dimensions[get_column_letter(j)].width = w

hoja_esperado("Esperado (con avance)",
              "CPM esperado con el escenario de «Avance (prueba)» (fecha de corte 15/03/2027), calculado por el script",
              ES2, EF2, LS2, LF2, FIN2, DUR2, ESTADO2)


# ---------------------------------------------------------------------------
# El mismo proyecto como lo exportarían MS Project (XML) y Primavera P6 (XER),
# con las fechas de «Esperado» como las que calculó el programa de origen.
# Sirven para probar la importación (P4) de punta a punta.
# ---------------------------------------------------------------------------
from xml.sax.saxutils import escape

TIPO_MSP = {"FS": 1, "SS": 3, "FF": 0, "SF": 2}
TIPO_P6 = {"FS": "PR_FS", "SS": "PR_SS", "FF": "PR_FF", "SF": "PR_SF"}
FIELD_TEXT1 = "188743731"


def iso(d, hora):
    return f"{d.isoformat()}T{hora}:00"


def generar_xml():
    x = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
         '<Project xmlns="http://schemas.microsoft.com/project">',
         "<SaveVersion>14</SaveVersion><Name>demo-cronograma.xml</Name><Title>Centro comunitario</Title>",
         f"<StartDate>{iso(INICIO, '08:00')}</StartDate><MinutesPerDay>480</MinutesPerDay><CalendarUID>1</CalendarUID>",
         "<ExtendedAttributes><ExtendedAttribute>"
         f"<FieldID>{FIELD_TEXT1}</FieldID><FieldName>Text1</FieldName><Alias>Categoría</Alias>"
         "</ExtendedAttribute></ExtendedAttributes>",
         "<Calendars><Calendar><UID>1</UID><Name>Estándar</Name><IsBaseCalendar>1</IsBaseCalendar><WeekDays>"]
    for dia in range(1, 8):  # 1 = domingo
        trabaja = 1 if dia in (2, 3, 4, 5, 6) else 0
        x.append(f"<WeekDay><DayType>{dia}</DayType><DayWorking>{trabaja}</DayWorking></WeekDay>")
    x.append("</WeekDays><Exceptions>")
    for f, nombre in FESTIVOS:
        x.append(f"<Exception><TimePeriod><FromDate>{iso(f, '00:00')}</FromDate><ToDate>{iso(f, '23:59')}</ToDate>"
                 f"</TimePeriod><Name>{escape(nombre)}</Name><DayWorking>0</DayWorking></Exception>")
    x.append("</Exceptions></Calendar></Calendars><Tasks>")
    x.append("<Task><UID>0</UID><ID>0</ID><Name>Centro comunitario</Name><OutlineLevel>0</OutlineLevel><Summary>1</Summary></Task>")
    for a in ACTIVIDADES:
        k, d = a[0], a[2]
        ini, fin = fecha_inicio(ES[k], d), fecha_fin(EF[k], d)
        partes = [f"<Task><UID>{k}</UID><ID>{k}</ID><Name>{escape(a[1])}</Name><OutlineLevel>1</OutlineLevel>",
                  f"<Summary>0</Summary><Milestone>{1 if d == 0 else 0}</Milestone>",
                  f"<Start>{iso(ini, '08:00')}</Start><Finish>{iso(fin, '17:00')}</Finish>",
                  f"<Duration>PT{int(d * 8)}H0M0S</Duration><DurationFormat>7</DurationFormat>"]
        if a[9]:
            partes.append(f"<ConstraintType>4</ConstraintType><ConstraintDate>{iso(a[9], '08:00')}</ConstraintDate>")
        else:
            partes.append("<ConstraintType>0</ConstraintType>")
        for p_, t_, r_ in leer_preds(a[3]):
            partes.append(f"<PredecessorLink><PredecessorUID>{p_}</PredecessorUID><Type>{TIPO_MSP[t_]}</Type>"
                          f"<CrossProject>0</CrossProject><LinkLag>{int(r_ * 4800)}</LinkLag><LagFormat>7</LagFormat></PredecessorLink>")
        partes.append(f"<ExtendedAttribute><FieldID>{FIELD_TEXT1}</FieldID><Value>{escape(a[7])}</Value></ExtendedAttribute>")
        partes.append("</Task>")
        x.append("".join(partes))
    x.append("</Tasks></Project>")
    with open("ejemplos/demo-cronograma.xml", "w", encoding="utf-8") as f:
        f.write("\n".join(x))


def generar_xer():
    fecha = lambda d, h: f"{d.isoformat()} {h}"
    dia_serie = lambda d: (d - dt.date(1899, 12, 30)).days
    # P6 guarda los saltos de línea de clndr_data como el carácter 127 (DEL).
    semana = []
    for dia in range(1, 8):
        if dia in (2, 3, 4, 5, 6):
            semana.append(f"    (0||{dia}()(\x7f\x7f      (0||0(s|08:00|f|12:00)())\x7f\x7f      (0||1(s|13:00|f|17:00)())))")
        else:
            semana.append(f"    (0||{dia}()())")
    excepciones = [f"    (0||{i}(d|{dia_serie(f)})())" for i, (f, _) in enumerate(FESTIVOS)]
    clndr = ("(0||CalendarData()(\x7f\x7f  (0||DaysOfWeek()(\x7f\x7f" + "\x7f\x7f".join(semana) + "))\x7f\x7f"
             "  (0||VIEW(ShowTotal|Y)())\x7f\x7f  (0||Exceptions()(\x7f\x7f" + "\x7f\x7f".join(excepciones) + "))))")
    L = ["ERMHDR\t19.12\t2027-01-15\tProject\tADMIN\tdbxDatabaseNoName\tProject Management\tUSD"]
    def tabla(nombre, campos, filas):
        L.append(f"%T\t{nombre}")
        L.append("%F\t" + "\t".join(campos))
        for f in filas:
            L.append("%R\t" + "\t".join(str(v) for v in f))
    tabla("CALENDAR", ["clndr_id", "default_flag", "clndr_name", "day_hr_cnt", "clndr_data"],
          [[1, "Y", "Estándar 5x8", 8, clndr]])
    tabla("PROJECT", ["proj_id", "proj_short_name", "plan_start_date", "last_recalc_date", "clndr_id"],
          [[100, "CENTRO-COM", fecha(INICIO, "08:00"), fecha(INICIO, "08:00"), 1]])
    cats = sorted({a[7] for a in ACTIVIDADES})
    wbs = {c: 200 + i for i, c in enumerate(cats)}
    tabla("PROJWBS", ["wbs_id", "proj_id", "proj_node_flag", "wbs_short_name", "wbs_name", "parent_wbs_id"],
          [[199, 100, "Y", "CENTRO-COM", "Centro comunitario", ""]] + [[wbs[c], 100, "N", c[:8].upper(), c, 199] for c in cats])
    filas = []
    for a in ACTIVIDADES:
        k, d = a[0], a[2]
        ini, fin = fecha_inicio(ES[k], d), fecha_fin(EF[k], d)
        tipo = "TT_Mile" if d == 0 and not a[3] else ("TT_FinMile" if d == 0 else "TT_Task")
        filas.append([1000 + int(k), 100, wbs[a[7]], 1, f"A{int(k) * 10}", a[1], tipo, "TK_NotStart",
                      d * 8, d * 8, fecha(ini, "08:00"), fecha(fin, "17:00"), "", "",
                      "CS_MSOA" if a[9] else "", fecha(a[9], "08:00") if a[9] else ""])
    tabla("TASK", ["task_id", "proj_id", "wbs_id", "clndr_id", "task_code", "task_name", "task_type", "status_code",
                   "target_drtn_hr_cnt", "remain_drtn_hr_cnt", "early_start_date", "early_end_date",
                   "act_start_date", "act_end_date", "cstr_type", "cstr_date"], filas)
    preds = []
    for a in ACTIVIDADES:
        for p_, t_, r_ in leer_preds(a[3]):
            preds.append([5000 + len(preds), 1000 + int(a[0]), 1000 + int(p_), 100, 100, TIPO_P6[t_], r_ * 8])
    tabla("TASKPRED", ["task_pred_id", "task_id", "pred_task_id", "proj_id", "pred_proj_id", "pred_type", "lag_hr_cnt"], preds)
    L.append("%E")
    with open("ejemplos/demo-cronograma.xer", "w", encoding="utf-8", newline="") as f:
        f.write("\r\n".join(L) + "\r\n")


generar_xml()
generar_xer()

wb.save("ejemplos/demo-cronograma.xlsx")
print(f"ejemplos/demo-cronograma.xlsx — fin {fecha_fin(FIN, 1)}, {FIN:g} días; con avance {fecha_fin(FIN2, 1)} ({FIN2:g}); "
      f"P50 {PCT[0.5] + 1} · P80 {PCT[0.8] + 1} · P90 {PCT[0.9] + 1} días; salud {SALUD}")
